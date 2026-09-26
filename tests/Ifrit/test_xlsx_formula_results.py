"""What a monster xlsx holds besides its values (Ifrit/IfritXlsx/xlsxmanager.py).

- Every stat formula is written with the result it computes. A spreadsheet that does not
  recalculate on opening (LibreOffice's default for an xlsx) shows that result - it used to be 0
  everywhere - and one that does recalculate 147 000 formulas of a 200-monster workbook takes ages.
- A drop-down reads its list from ref_data row $2, absolute: a relative first row moved down with
  every cell of a multi-cell drop-down, and Draw 4 no longer offered the first three spells.
- Cells with the same rule share one data validation, not one each.
"""
import contextlib
import io
import pathlib
import re
import zipfile

import openpyxl
import pytest

from Ifrit.IfritStat.ifritstatwidget import StatCurvePlot
from Ifrit.IfritXlsx.xlsxmanager import (DEFAULT_MONSTER_LVL, EXCEL_DIV_ZERO, cell_result, ref_data_list_source,
                                         stat_impacts, stat_total)

PROJECT_ROOT = pathlib.Path(__file__).parent.parent.parent
BATTLE_FILE = "extracted_files/battle/c0m071.dat"
STAT_NAMES = ['hp', 'str', 'vit', 'mag', 'spr', 'spd', 'eva']


@pytest.mark.parametrize("stat_name", STAT_NAMES)
@pytest.mark.parametrize("stat_bytes", [[1, 2, 3, 4], [30, 40, 0, 1], [255, 255, 255, 255], [7, 1, 50, 130]])
def test_the_formula_results_are_the_game_curve(stat_name, stat_bytes):
    """With no 0 divisor, the xlsx formula is the curve the stat editor draws - the game's."""
    for level in (1, 10, 55, 100):
        assert stat_total(stat_impacts(stat_name, stat_bytes, level), stat_name) == StatCurvePlot._stat_value(stat_name, stat_bytes, level)


@pytest.mark.parametrize("stat_name", ['str', 'vit'])
def test_a_formula_dividing_by_a_0_byte_is_an_error(stat_name):
    impacts = stat_impacts(stat_name, [5, 0, 5, 5], 10)
    assert impacts[1] is None and impacts[3] is not None
    assert cell_result(impacts[1]) == EXCEL_DIV_ZERO
    assert cell_result(stat_total(impacts)) == EXCEL_DIV_ZERO


def test_a_drop_down_source_is_absolute_on_both_rows():
    assert ref_data_list_source(2, 97) == "=ref_data!$C$2:$C$98"


def _write_monster(path, with_stat_graph=True):
    """c0m071 written as the editor writes it, and the monster it was written from."""
    from FF8GameData.dat.monsteranalyser import MonsterAnalyser
    from FF8GameData.gamedata import GameData
    from Ifrit.IfritAI.AICompiler.AIDecompiler import AIDecompiler
    from Ifrit.IfritXlsx.xlsxmanager import DatToXlsx

    game_data = GameData(str(PROJECT_ROOT / "FF8GameData"))
    game_data.load_all()
    monster = MonsterAnalyser(game_data)
    with contextlib.redirect_stdout(io.StringIO()):
        monster.load_file_data(str(PROJECT_ROOT / BATTLE_FILE), game_data)
        monster.analyse_loaded_data(game_data, AIDecompiler(game_data))
        writer = DatToXlsx()
        writer.create_file(str(path), with_stat_graph)
        writer.export_to_xlsx(monster, "c0m071.dat", game_data, analyse_ai=False)
        writer.create_ref_data(game_data)
        writer.close_file()
    return path, monster


@pytest.fixture(scope="module")
def written(tmp_path_factory):
    return _write_monster(tmp_path_factory.mktemp("xlsx") / "monster.xlsx")


@pytest.mark.ff8data(BATTLE_FILE)
def test_a_workbook_without_the_stat_graph_holds_everything_else(written, tmp_path):
    """The to-edit workbook of a mod leaves out the chart and its level table; the rest is the same."""
    with_graph, _monster = written
    without_graph, _monster = _write_monster(tmp_path / "no_graph.xlsx", with_stat_graph=False)
    with zipfile.ZipFile(without_graph) as workbook:
        assert not [name for name in workbook.namelist() if "chart" in name or "drawing" in name]

    def cells(path):
        sheet = openpyxl.load_workbook(path).worksheets[0]
        return {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}

    graph_cells = cells(with_graph)
    level_table = {coordinate for coordinate in graph_cells
                   if openpyxl.utils.cell.coordinate_to_tuple(coordinate)[0] >= 34
                   and 5 <= openpyxl.utils.cell.coordinate_to_tuple(coordinate)[1] <= 12}   # E34:L134
    assert len(level_table) == 8 * 101   # the levels and 7 stats, each with its title
    assert cells(without_graph) == {coordinate: value for coordinate, value in graph_cells.items()
                                    if coordinate not in level_table}


@pytest.mark.ff8data(BATTLE_FILE)
def test_every_stat_formula_holds_its_result(written):
    path, monster = written
    formulas = openpyxl.load_workbook(path).worksheets[0]
    results = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    for row, stat_name in enumerate(STAT_NAMES, start=2):
        stat_bytes = monster.info_stat_data[stat_name]
        assert [formulas.cell(row, column).value for column in range(5, 9)] == stat_bytes   # E..H: the 4 bytes
        impacts = stat_impacts(stat_name, stat_bytes, DEFAULT_MONSTER_LVL)
        assert [results.cell(row, column).value for column in range(9, 13)] == [cell_result(x) for x in impacts]   # I..L
        assert results.cell(row, 13).value == cell_result(stat_total(impacts, stat_name))   # M: total
        # The level table: level L on row 34 + L, one column per stat from F
        for level in (1, 50, 100):
            expected = StatCurvePlot._stat_value(stat_name, stat_bytes, level)
            if stat_name == 'hp':
                expected /= 100
            assert str(formulas.cell(34 + level, 4 + row).value).startswith('=')
            assert results.cell(34 + level, 4 + row).value == expected


@pytest.mark.ff8data(BATTLE_FILE)
def test_the_drop_downs_share_a_rule_and_read_their_whole_list(written):
    path, _monster = written
    with zipfile.ZipFile(path) as workbook:
        sheet = workbook.read("xl/worksheets/sheet1.xml").decode("utf8")
    rules = re.findall(r'<dataValidation [^>]*sqref="([^"]+)"[^>]*>.*?(?:<formula1>(.*?)</formula1>|</dataValidation>)', sheet, re.S)
    sources = [source for _cells, source in rules if source.startswith("ref_data")]
    assert sources and all(re.fullmatch(r"ref_data!\$[A-Z]\$2:\$[A-Z]\$\d+", source) for source in sources)
    assert len(sources) == len(set(sources))   # one rule per list, not one per cell
    draw_rule = [cells.split() for cells, source in rules if source.startswith("ref_data!$C$")][0]
    assert {"S2:S5", "U2:U5", "W2:W5"} <= set(draw_rule)   # Draw 1-4 of the three levels
