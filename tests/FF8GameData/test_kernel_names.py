"""The names a kernel.bin owns (FF8GameData/kernelnames.py) and showing them in a monster xlsx.

A mod renames things in its kernel.bin - Cronos renames a few spells and items - and the tools
show the vanilla names from the json until those renames are read back out of the file. What is
pinned here:

- the kernel sections line up with the json lists (the ids are the same), which is the whole
  premise: read a name at kernel index i, it is the name of id i of that list,
- only renames are taken, never the whole kernel list: the json labels the entries the game file
  leaves unnamed ("Not defined", "Physical attack"...) and those labels stay,
- applying them changes what every tool shows, and a workbook can be refreshed to match without
  losing anything it holds - a name in it is decoration, the id in front of it is the value.
"""
import contextlib
import io
import pathlib
import re
import zipfile

import pytest

from FF8GameData import kernelnames
from FF8GameData.gamedata import GameData
from Ifrit.IfritXlsx import xlsxnames

PROJECT_ROOT = pathlib.Path(__file__).parent.parent.parent
KERNEL = "extracted_files/main/kernel.bin"
BATTLE_FILE = "extracted_files/battle/c0m071.dat"


@pytest.fixture(scope="module")
def game_data():
    data = GameData(str(PROJECT_ROOT / "FF8GameData"))
    data.load_all()
    return data


@pytest.mark.ff8data(KERNEL)
def test_the_kernel_sections_hold_the_names_of_the_json_lists(game_data):
    """Every name the vanilla kernel.bin has is the json's name for that id - if this ever fails,
    a list is being read at the wrong ids and every rename would land on the wrong entry."""
    kernel_names = kernelnames.read_names(game_data, PROJECT_ROOT / KERNEL)
    json_names = kernelnames.json_names(game_data)
    assert set(kernel_names) == {"magic", "item", "enemy_ability"}
    assert len(kernel_names["magic"]) == 56          # 57 spells, the first one is unnamed
    assert len(kernel_names["item"]) == 198          # the 33 battle items, then all the others
    for list_name, names in kernel_names.items():
        differing = {id_: (name, json_names[list_name].get(id_))
                     for id_, name in names.items() if json_names[list_name].get(id_) != name}
        # A handful of entries are spelled differently by the shipped json (Ray-Bomb / Ray Bomb)
        assert len(differing) < 20, f"{list_name}: {differing}"
    assert kernel_names["magic"][43] == json_names["magic"][43] == "Death"
    assert kernel_names["item"][144] == json_names["item"][144] == "Arctic Wind"


@pytest.mark.ff8data(KERNEL)
def test_a_kernel_compared_with_itself_renames_nothing(game_data):
    assert kernelnames.name_changes(game_data, PROJECT_ROOT / KERNEL,
                                    vanilla_kernel_file=PROJECT_ROOT / KERNEL) == {}


@pytest.mark.ff8data(KERNEL)
def test_the_unnamed_entries_keep_the_labels_the_json_gives_them(game_data):
    """The kernel leaves an enemy attack unnamed when the game never shows it; the json calls it
    "Not defined" or "Physical attack". Reading names from a file must not wipe that."""
    changes = kernelnames.name_changes(game_data, PROJECT_ROOT / KERNEL)
    assert not any(0 in names for names in changes.values())
    kernelnames.apply_names(game_data, changes)
    abilities = {entry["id"]: entry["name"] for entry in game_data.enemy_abilities_data_json["abilities"]}
    assert abilities[0] == "Not defined" and abilities[2] == "Physical attack"
    game_data.load_names()   # Back to the shipped names for the tests that follow


def test_applying_renames_changes_what_every_tool_shows(game_data):
    changes = {"magic": {43: "Reaper"}, "item": {144: "Stone Skin"}, "enemy_ability": {16: "Arm Machine Gun"}}
    assert kernelnames.apply_names(game_data, changes) == 3
    assert kernelnames.json_names(game_data)["magic"][43] == "Reaper"
    assert kernelnames.json_names(game_data)["item"][144] == "Stone Skin"
    assert kernelnames.json_names(game_data)["enemy_ability"][16] == "Arm Machine Gun"
    game_data.load_names()   # The vanilla names are what a GameData holds unless asked otherwise
    assert kernelnames.json_names(game_data)["magic"][43] == "Death"


def test_a_renames_file_survives_being_written_and_read(game_data, tmp_path):
    changes = {"magic": {43: "Reaper", 54: "Earthra"}, "item": {144: "Stone Skin"}}
    path = tmp_path / "names_cronos.json"
    kernelnames.write_changes_file(path, changes, "written by a test")
    assert kernelnames.read_changes_file(path) == changes   # the comment line is not a name list


def test_game_data_loads_a_mods_names_and_goes_back_to_the_vanilla_ones(tmp_path):
    """GameData.load_names is what the Ifrit Cronos checkbox and the Cronos tool both call."""
    data = GameData(str(PROJECT_ROOT / "FF8GameData"))
    data.load_all()
    data.load_names("names_cronos.json")
    assert kernelnames.json_names(data)["magic"][43] == "Reaper"
    data.load_names()
    assert kernelnames.json_names(data)["magic"][43] == "Death"


# ---------------------------------------------------------------------------------------------
# Showing them in a workbook
# ---------------------------------------------------------------------------------------------

def test_only_an_unambiguous_name_is_renamed():
    """Every list writes "<id>:<name>", so a text is only renamed when one list could have written
    it: two lists with the same id AND the same name are reported instead."""
    old = {"magic": {43: "Death"}, "item": {43: "Death", 18: "Hero"}}
    new = {"magic": {43: "Reaper"}, "item": {43: "Reaper", 18: "Grenade"}}
    assert xlsxnames.name_renames(old, new) == {"43:Death": "43:Reaper", "18:Hero": "18:Grenade"}
    assert xlsxnames.ambiguous_renames(old, new) == {}

    new_disagreeing = {"magic": {43: "Reaper"}, "item": {43: "Doom"}}
    assert "43:Death" not in xlsxnames.name_renames(old, new_disagreeing)
    assert xlsxnames.ambiguous_renames(old, new_disagreeing) == {
        "43:Death": {"magic": "43:Reaper", "item": "43:Doom"}}


def _monster_workbook(game_data, tmp_path, monster_file=None):
    """One monster written as the editor writes it: its sheet, and the ref_data sheet holding the
    lists every drop-down of it points at."""
    from FF8GameData.dat.monsteranalyser import MonsterAnalyser
    from Ifrit.IfritAI.AICompiler.AIDecompiler import AIDecompiler
    from Ifrit.IfritXlsx.xlsxmanager import DatToXlsx

    monster = MonsterAnalyser(game_data)
    path = tmp_path / "monster.xlsx"
    with contextlib.redirect_stdout(io.StringIO()):
        monster.load_file_data(str(monster_file or PROJECT_ROOT / BATTLE_FILE), game_data)
        monster.analyse_loaded_data(game_data, AIDecompiler(game_data))
        writer = DatToXlsx()
        writer.create_file(str(path))
        writer.export_to_xlsx(monster, "c0m071.dat", game_data, analyse_ai=False)
        writer.create_ref_data(game_data)
        writer.close_file()
    return path


@pytest.mark.ff8data(BATTLE_FILE)
def test_renaming_a_workbook_keeps_everything_it_holds(game_data, tmp_path):
    """The names shown change; the workbook itself - its charts, its drop-downs, the values it
    holds - is the same file with one part rewritten."""
    from FF8GameData.dat.monsteranalyser import MonsterAnalyser
    from Ifrit.IfritAI.AICompiler.AIDecompiler import AIDecompiler

    monster = MonsterAnalyser(game_data)
    with contextlib.redirect_stdout(io.StringIO()):
        monster.load_file_data(str(PROJECT_ROOT / BATTLE_FILE), game_data)
        monster.analyse_loaded_data(game_data, AIDecompiler(game_data))
    path = _monster_workbook(game_data, tmp_path)
    before = {info.filename: zipfile.ZipFile(path).read(info.filename)
              for info in zipfile.ZipFile(path).infolist()}

    renames = xlsxnames.name_renames({"magic": {43: "Death"}}, {"magic": {43: "Reaper"}})
    assert xlsxnames.rename_in_workbook(path, renames) == {"43:Death": 1}

    after = {info.filename: zipfile.ZipFile(path).read(info.filename)
             for info in zipfile.ZipFile(path).infolist()}
    assert list(after) == list(before)
    assert [name for name in after if after[name] != before[name]] == ["xl/sharedStrings.xml"]
    assert b">43:Reaper<" in after["xl/sharedStrings.xml"]

    # And it still reads as the same monster: a name is shown, an id is what is read back
    from Ifrit.IfritXlsx import xlsxmanager
    reader = xlsxmanager.XlsxToDat()
    reader.load_file(str(path))
    sheet = next(sheet for sheet in reader.workbook if sheet.title != xlsxmanager.REF_DATA_SHEET_TITLE)
    with contextlib.redirect_stdout(io.StringIO()):
        stats = reader.get_stat_info(sheet, game_data)
    reader.close_file()
    assert stats["high_lvl_mag"] == monster.info_stat_data["high_lvl_mag"]     # the spells drawn
    assert stats["high_lvl_drop"] == monster.info_stat_data["high_lvl_drop"]
    assert stats["abilities_high"] == monster.info_stat_data["abilities_high"]


@pytest.mark.ff8data(BATTLE_FILE)
def test_a_workbook_without_those_names_is_left_alone(game_data, tmp_path):
    path = tmp_path / "empty.xlsx"
    path.write_bytes(b"not a workbook")
    assert xlsxnames.rename_in_workbook(path, {}) == {}
    assert path.read_bytes() == b"not a workbook"


# ---------------------------------------------------------------------------------------------
# A mod that ADDS spells: more of them to offer
# ---------------------------------------------------------------------------------------------

def _kernel_with_more_magic(game_data, source, destination, added_names: dict):
    """A kernel.bin with a longer magic section, the way SolomonRing's "+ Add entry" grows it:
    blank entries appended, and the ids 64-95 the game keeps for the GF summons skipped over."""
    from ShumiTranslator.model.kernel.kernelmanager import KernelManager

    kernel_manager = KernelManager(game_data)
    kernel_manager.load_file(str(source))
    by_id = {section.id: section for section in kernel_manager.section_list if section}
    magic, texts = by_id[2], by_id[33]
    for _ in range(max(added_names) + 1 - len(magic.get_subsection_list())):
        magic.append_blank_subsection()
        texts.add_text(bytearray([0x00]))
        texts.add_text(bytearray([0x00]))
    for id_, name in added_names.items():
        texts.get_text_list()[id_ * 2].set_str(name)
    kernel_manager.save_file(str(destination))
    return destination


@pytest.mark.ff8data(KERNEL)
def test_spells_a_mod_adds_are_read_and_offered(game_data, tmp_path):
    """Ids 57 to 63 are free, 64 to 95 belong to the GF summons and 96 up is free again. A spell
    added in a free id is a change like a rename, and one past the end of the list is added to it."""
    grown = _kernel_with_more_magic(game_data, PROJECT_ROOT / KERNEL, tmp_path / "grown.bin",
                                    {57: "Blast Wave", 97: "Meteor Rain"})
    changes = kernelnames.name_changes(game_data, grown, vanilla_kernel_file=PROJECT_ROOT / KERNEL)
    assert changes == {"magic": {57: "Blast Wave", 97: "Meteor Rain"}}

    spells_before = len(game_data.magic_data_json["magic"])
    assert kernelnames.apply_names(game_data, changes) == 2
    spells = {entry["id"]: entry["name"] for entry in game_data.magic_data_json["magic"]}
    assert spells[57] == "Blast Wave"                      # was "Unknown 0x39", a free id
    assert spells[97] == "Meteor Rain"                     # past the end: the list is longer now
    assert len(game_data.magic_data_json["magic"]) == spells_before + 1
    assert [entry["id"] for entry in game_data.magic_data_json["magic"]] == sorted(spells)
    # The rows a grown section holds for the GF summons are padding, not names
    assert spells[64] == "Thunder Storm (Quezacotl)"
    game_data.load_names()


@pytest.mark.ff8data(BATTLE_FILE)
def test_a_workbook_offers_the_spells_added_to_it(game_data, tmp_path):
    """The names are picked from a column of the ref_data sheet, so offering more of them is
    writing them there and stretching every drop-down that reads it."""
    from openpyxl import load_workbook

    path = _monster_workbook(game_data, tmp_path)
    offered = xlsxnames.list_length_in_workbook(path, "magic")
    assert offered == len(game_data.magic_data_json["magic"])

    texts = [f"{entry['id']}:{entry['name']}" for entry in game_data.magic_data_json["magic"]]
    assert xlsxnames.grow_list_in_workbook(path, "magic", texts) == 0   # nothing new to offer
    texts += ["97:Meteor Rain", "98:Blast Wave"]
    assert xlsxnames.grow_list_in_workbook(path, "magic", texts) == 2
    assert xlsxnames.list_length_in_workbook(path, "magic") == offered + 2

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    reference = workbook["ref_data"]
    column = xlsxnames.REF_DATA_COLUMN["magic"] + 1
    assert [reference.cell(row=offered + row, column=column).value for row in (1, 2, 3)] == \
           [texts[-3], "97:Meteor Rain", "98:Blast Wave"]
    workbook.close()

    with zipfile.ZipFile(path) as opened:
        sheet = opened.read("xl/worksheets/sheet1.xml").decode("utf8")
    assert f"ref_data!$C2:$C${offered + 3}" in sheet   # the drop-down reaches the new rows


@pytest.mark.ff8data(BATTLE_FILE)
def test_a_workbook_is_written_with_the_names_in_use(game_data, tmp_path):
    """Every workbook this repository writes - the no-edit one a Cronos build produces included -
    shows the names loaded at the time, so a mod's own names need no pass over the file."""
    game_data.load_names("names_cronos.json")
    try:
        path = _monster_workbook(game_data, tmp_path)
        with zipfile.ZipFile(path) as workbook:
            texts = set(re.findall(r"<t[^>]*>([^<]*)</t>",
                                   workbook.read("xl/sharedStrings.xml").decode("utf8")))
        assert {"43:Reaper", "144:Stone Skin", "16:Arm Machine Gun"} <= texts
        assert not {"43:Death", "144:Arctic Wind", "16:Arm Slash"} & texts
    finally:
        game_data.load_names()


@pytest.mark.ff8data(BATTLE_FILE)
def test_the_bit_flags_are_relabelled_without_moving_a_value(game_data, tmp_path):
    """A bit is its position; its name is only what the tools call it, and names keep being found.
    Relabelling a workbook must therefore move nothing - not a value, not a row."""
    from FF8GameData.monsterdata import AIData
    from openpyxl import load_workbook
    from Ifrit.IfritXlsx.xlsxmanager import COL_MISC, ROW_BYTE_FLAG

    path = _monster_workbook(game_data, tmp_path)
    assert xlsxnames.refresh_byte_flag_labels(path) == {}    # just written: already current

    # A workbook written when byte 2's bits were still "byte2_zz1", "byte2_unused_3"...
    found_since = {AIData.BYTE_FLAG_VALUES["byte_flag_2"][index]: name for index, name in
                   enumerate(["byte2_zz1", "byte2_zz2", "byte2_unused_3", "byte2_unused_4", "byte2_unused_5"])}
    assert xlsxnames.rename_in_workbook(path, found_since) == {name: 1 for name in found_since}

    before = load_workbook(path, data_only=True, keep_links=False)
    sheet = before[[name for name in before.sheetnames if name != "ref_data"][0]]
    values = [sheet.cell(row=ROW_BYTE_FLAG + 1 + bit, column=COL_MISC + 2).value for bit in range(32)]
    before.close()

    assert xlsxnames.refresh_byte_flag_labels(path) == {was: now for now, was in found_since.items()}

    after = load_workbook(path, data_only=True, keep_links=False)
    sheet = after[[name for name in after.sheetnames if name != "ref_data"][0]]
    assert [sheet.cell(row=ROW_BYTE_FLAG + 1 + bit, column=COL_MISC + 1).value for bit in range(32)] == \
           [label for flag in AIData.BYTE_FLAG_LIST for label in AIData.BYTE_FLAG_VALUES[flag]]
    assert [sheet.cell(row=ROW_BYTE_FLAG + 1 + bit, column=COL_MISC + 2).value for bit in range(32)] == values
    after.close()
